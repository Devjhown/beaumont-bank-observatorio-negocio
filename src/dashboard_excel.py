"""
Gera um dashboard em Excel (reports/dashboard.xlsx) com os principais KPIs
da carteira de empréstimos e um gráfico de status, usando openpyxl.
"""
import os
import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

BASE_DIR = os.path.join(os.path.dirname(__file__), '..')
DB_PATH = os.path.join(BASE_DIR, 'data', 'beaumont_bank.db')
OUT_PATH = os.path.join(BASE_DIR, 'reports', 'dashboard.xlsx')

conn = sqlite3.connect(DB_PATH)
loan = pd.read_sql('SELECT * FROM loan', conn)
conn.close()

loan['inadimplente'] = loan['status'].isin(['B', 'D'])

kpis = {
    'Total de empréstimos': len(loan),
    'Total emprestado (CZK)': round(loan['amount'].sum(), 2),
    'Ticket médio (CZK)': round(loan['amount'].mean(), 2),
    'Duração média (meses)': round(loan['duration'].mean(), 1),
    'Taxa de inadimplência': f"{loan['inadimplente'].mean():.1%}",
}
status_counts = loan['status'].value_counts().sort_index()

wb = Workbook()

# --- Aba de KPIs ---
ws_kpi = wb.active
ws_kpi.title = 'KPIs'
ws_kpi['A1'] = 'Observatório de Negócio — Beaumont Bank'
ws_kpi['A1'].font = Font(size=14, bold=True)
ws_kpi.merge_cells('A1:B1')

header_fill = PatternFill('solid', fgColor='2B6CB0')
ws_kpi['A3'] = 'Indicador'
ws_kpi['B3'] = 'Valor'
for cell in ('A3', 'B3'):
    ws_kpi[cell].font = Font(bold=True, color='FFFFFF')
    ws_kpi[cell].fill = header_fill

row = 4
for nome, valor in kpis.items():
    ws_kpi[f'A{row}'] = nome
    ws_kpi[f'B{row}'] = valor
    row += 1

ws_kpi.column_dimensions['A'].width = 30
ws_kpi.column_dimensions['B'].width = 20

# --- Aba de dados + gráfico ---
ws_data = wb.create_sheet('Status dos Empréstimos')
ws_data['A1'] = 'Status'
ws_data['B1'] = 'Quantidade'
for i, (status, qtd) in enumerate(status_counts.items(), start=2):
    ws_data[f'A{i}'] = status
    ws_data[f'B{i}'] = int(qtd)

chart = BarChart()
chart.title = 'Empréstimos por status'
chart.x_axis.title = 'Status'
chart.y_axis.title = 'Quantidade'
data_ref = Reference(ws_data, min_col=2, min_row=1, max_row=1 + len(status_counts))
cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=1 + len(status_counts))
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
ws_data.add_chart(chart, 'D2')

wb.save(OUT_PATH)
print(f'Dashboard salvo em: {OUT_PATH}')
